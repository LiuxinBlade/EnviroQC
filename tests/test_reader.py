from __future__ import annotations

import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from enviroqc import (
    CorruptFileError,
    EmptyFileError,
    EncodingError,
    FileNotFoundInputError,
    HeaderError,
    InputFileError,
    InputOptions,
    UnsupportedFormatError,
    read_table,
)

from .conftest import sha256


def rows(table):
    return [(r.row, r.values) for r in table.records]


def test_csv_rows_are_numbered_by_source_line(write_csv):
    p = write_csv("a,b\n1,2\n3,4\n")
    t = read_table(p)
    assert t.header == ("a", "b")
    assert t.header_row == 1
    assert rows(t) == [(2, ("1", "2")), (3, ("3", "4"))]
    assert t.format == "csv" and t.sheet is None and t.file == "data.csv"


def test_csv_row_numbers_survive_multiline_quoted_fields(write_csv):
    p = write_csv('a,b\n1,"line one\nline two\nline three"\n2,x\n3,y\n')
    t = read_table(p)
    assert [r.row for r in t.records] == [2, 5, 6]
    assert t.records[0].values[1] == "line one\nline two\nline three"


def test_csv_crlf_line_endings(write_csv):
    p = write_csv('a,b\r\n1,2\r\n"x\r\ny",3\r\n4,5\r\n')
    assert [r.row for r in read_table(p).records] == [2, 3, 5]


def test_csv_blank_lines_are_skipped_but_counted(write_csv):
    p = write_csv("\n\na,b\n1,2\n\n , \n3,4\n")
    t = read_table(p)
    assert t.header_row == 3
    assert [r.row for r in t.records] == [4, 7]


def test_csv_header_row_skips_title_lines(write_csv):
    p = write_csv('Station report "draft\nexported 2024\na,b\n1,2\n')
    t = read_table(p, InputOptions(header_row=3))
    assert t.header == ("a", "b")
    assert t.header_row == 3
    assert rows(t) == [(4, ("1", "2"))]


def test_csv_utf8_bom_is_removed_from_header(write_csv):
    p = write_csv("﻿station,值\nST1,1\n")
    assert read_table(p).header == ("station", "值")


def test_csv_header_whitespace_and_trailing_empty_columns(write_csv):
    p = write_csv(" a , b ,,\n1,2,,\n")
    t = read_table(p)
    assert t.header == ("a", "b")
    assert t.records[0].values == ("1", "2")


def test_csv_extra_non_empty_cells_are_kept_for_malformed_row_check(write_csv):
    t = read_table(write_csv("a,b\n1,2,3\n4\n"))
    assert t.records[0].values == ("1", "2", "3")
    assert t.records[1].values == ("4",)


def test_csv_custom_delimiter_and_tsv(write_csv):
    p = write_csv("a;b\n1,5;2\n")
    assert read_table(p, InputOptions(delimiter=";")).records[0].values == ("1,5", "2")
    tsv = write_csv("a\tb\n1\t2\n", name="data.tsv")
    assert read_table(tsv).records[0].values == ("1", "2")


def test_csv_gbk_requires_encoding(write_csv):
    p = write_csv("站点,温度\n北京,1.5\n", encoding="gbk")
    with pytest.raises(EncodingError) as excinfo:
        read_table(p)
    assert "line 1" in str(excinfo.value)
    assert "gbk" in str(excinfo.value)
    t = read_table(p, InputOptions(encoding="gbk"))
    assert t.header == ("站点", "温度")
    assert t.records[0].values == ("北京", "1.5")


def test_csv_encoding_error_reports_correct_line(write_csv, tmp_path):
    p = tmp_path / "bad.csv"
    p.write_bytes(b"a,b\n1,2\n3,\xff\n")
    with pytest.raises(EncodingError, match="line 3"):
        read_table(p)


def test_csv_with_nul_bytes_is_rejected(tmp_path):
    p = tmp_path / "utf16.csv"
    p.write_bytes("a,b\n1,2\n".encode("utf-16-le"))
    with pytest.raises(EncodingError, match="NUL"):
        read_table(p)


def test_csv_unbalanced_quote_is_reported(write_csv):
    with pytest.raises(CorruptFileError, match="Malformed CSV"):
        read_table(write_csv('a,b\n1,"oops\n2,3\n'))


@pytest.mark.parametrize("content", ["", "   \n\n  ", "\n"])
def test_empty_file(write_csv, content):
    with pytest.raises(EmptyFileError):
        read_table(write_csv(content))


def test_header_only_file_has_no_records(write_csv):
    t = read_table(write_csv("a,b\n"))
    assert t.header == ("a", "b") and t.records == ()


def test_header_row_beyond_end_of_file(write_csv):
    with pytest.raises(EmptyFileError, match="No header row"):
        read_table(write_csv("a,b\n1,2\n"), InputOptions(header_row=10))


def test_duplicate_header_names(write_csv):
    with pytest.raises(HeaderError, match=r"duplicate column name.*'a'"):
        read_table(write_csv("a,b,a\n1,2,3\n"))


def test_blank_header_name_in_the_middle(write_csv):
    with pytest.raises(HeaderError, match="position"):
        read_table(write_csv("a,,c\n1,2,3\n"))


def test_missing_file_and_directory(tmp_path):
    with pytest.raises(FileNotFoundInputError, match="not found"):
        read_table(tmp_path / "missing.csv")
    with pytest.raises(FileNotFoundInputError, match="not a regular file"):
        read_table(tmp_path)


@pytest.mark.parametrize("name", ["data.xls", "data.json", "data"])
def test_unsupported_formats(tmp_path, name):
    p = tmp_path / name
    p.write_text("a,b\n1,2\n")
    with pytest.raises(UnsupportedFormatError):
        read_table(p)


def test_reading_never_modifies_the_file(write_csv):
    p = write_csv("a,b\n1,2\n")
    before = (sha256(p), p.stat().st_mtime_ns)
    read_table(p)
    assert (sha256(p), p.stat().st_mtime_ns) == before


# ---------------------------------------------------------------- XLSX


def make_xlsx(path: Path, sheets: dict[str, dict[tuple[int, int], object]]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for title, cells in sheets.items():
        ws = wb.create_sheet(title)
        for (r, c), v in cells.items():
            ws.cell(row=r, column=c, value=v)
    wb.save(path)
    return path


def test_xlsx_rows_use_excel_numbers_with_gaps(tmp_path):
    p = make_xlsx(
        tmp_path / "d.xlsx",
        {
            "S": {
                (1, 1): "a",
                (1, 2): "b",
                (2, 1): 1,
                (2, 2): 2.5,
                (5, 1): datetime(2024, 1, 1, 3),
                (6, 2): "x",
            }
        },
    )
    t = read_table(p)
    assert t.sheet == "S" and t.format == "xlsx"
    assert rows(t) == [
        (2, (1, 2.5)),
        (5, (datetime(2024, 1, 1, 3), None)),
        (6, (None, "x")),
    ]


def test_xlsx_first_used_row_is_found_and_blank_leading_column_is_reported(tmp_path):
    p = make_xlsx(tmp_path / "d.xlsx", {"S": {(3, 2): "a", (4, 2): 7}})
    # Rows 1-2 are empty and skipped; column A has no header -> clear error at row 3.
    with pytest.raises(HeaderError, match=r"Header row 3 .* position\(s\) 1"):
        read_table(p)


def test_xlsx_header_not_on_first_row(tmp_path):
    p = make_xlsx(
        tmp_path / "d.xlsx",
        {"S": {(1, 1): "Title", (3, 1): "a", (3, 2): "b", (4, 1): 1, (4, 2): 2}},
    )
    t = read_table(p, InputOptions(header_row=3))
    assert t.header == ("a", "b") and t.header_row == 3
    assert rows(t) == [(4, (1, 2))]


def test_xlsx_sheet_selection(tmp_path):
    p = make_xlsx(
        tmp_path / "d.xlsx",
        {"first": {(1, 1): "x", (2, 1): 1}, "second": {(1, 1): "y", (2, 1): 2, (3, 1): 3}},
    )
    assert read_table(p).header == ("x",)
    t = read_table(p, InputOptions(sheet="second"))
    assert t.sheet == "second" and t.header == ("y",) and len(t.records) == 2
    with pytest.raises(InputFileError, match="'third' not found") as excinfo:
        read_table(p, InputOptions(sheet="third"))
    assert "first, second" in str(excinfo.value)


def test_xlsx_empty_sheet(tmp_path):
    p = make_xlsx(tmp_path / "d.xlsx", {"S": {}})
    with pytest.raises(EmptyFileError, match="No header row"):
        read_table(p)


def test_xlsx_not_a_zip(tmp_path):
    p = tmp_path / "fake.xlsx"
    p.write_text("a,b\n1,2\n")
    with pytest.raises(CorruptFileError, match="not a ZIP"):
        read_table(p)


def test_xlsx_zip_without_workbook(tmp_path):
    p = tmp_path / "broken.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("hello.txt", "not excel")
    with pytest.raises(CorruptFileError, match="Cannot open"):
        read_table(p)


def test_xlsx_truncated_file(tmp_path):
    good = make_xlsx(tmp_path / "good.xlsx", {"S": {(1, 1): "a", (2, 1): 1}})
    bad = tmp_path / "truncated.xlsx"
    bad.write_bytes(good.read_bytes()[:-200])
    with pytest.raises(CorruptFileError):
        read_table(bad)
